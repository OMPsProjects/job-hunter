"""Resume intake: read a PDF / DOCX / TXT / MD resume, and load candidate.yaml.

candidate.yaml is the structured version of a resume that ranking and tailoring use.
`python job_hunter.py resume <file>` extracts the text and drafts the easy fields
(contact info, years of experience). Claude (or `--llm`) fills in titles, levels and skills.
See candidate.example.yaml.
"""
import datetime as dt
import re
import sys
from pathlib import Path

import yaml

LEVELS = ["intern", "entry", "mid", "senior", "manager", "director", "vp", "executive"]
LEVEL_INDEX = {name: i for i, name in enumerate(LEVELS)}

# Ordered most-specific first; the first pattern that hits the title wins.
_TITLE_LEVELS = [
    (r"\bintern(ship)?\b|\bsummer\b|\bco-?op\b", "intern"),
    (r"\bchief\b|\bc[eftoi]o\b|(?<!vice )(?<!vice-)\bpresident\b|\bpartner\b|\bfounder\b", "executive"),
    (r"\bevp\b|\bsvp\b|senior vice president|executive vice president|managing director", "vp"),
    (r"\bvice president\b|\bvp\b", "vp"),
    (r"\bassociate director\b|\bdirector\b|\bhead of\b", "director"),
    (r"\bmanager\b|\bteam lead\b|\blead\b|\bprincipal\b|\bsupervisor\b", "manager"),
    (r"\bsenior\b|\bsr\.?\b|\bstaff\b|\biii\b|\biv\b", "senior"),
    (r"\bjunior\b|\bjr\.?\b|\bentry\b|\bassistant\b|\bapprentice\b|\btrainee\b", "entry"),
    (r"\banalyst\b|\bassociate\b|\bspecialist\b|\bcoordinator\b|\bengineer\b|\bunderwriter\b|\bii\b", "mid"),
]


def title_level(title):
    t = str(title or "").lower()
    for rx, level in _TITLE_LEVELS:
        if re.search(rx, t):
            return LEVEL_INDEX[level]
    return None


def years_required(text):
    """Smallest 'N+ years ... experience' requirement in a job description, or None."""
    t = str(text or "").lower()
    found = []
    for m in re.finditer(r"(\d{1,2})\s*(?:\+|plus)?\s*(?:-|–|to)?\s*(\d{1,2})?\s*\+?\s*(?:years?|yrs?)\b", t):
        window = t[m.end(): m.end() + 80] + t[max(0, m.start() - 40): m.start()]
        if "experience" in window or "exp" in window:
            n = int(m.group(1))
            if 0 < n <= 30:
                found.append(n)
    return min(found) if found else None


# ----------------------------- reading files --------------------------------

def read_resume_text(path: Path) -> str:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            sys.exit("pip install pypdf  (needed to read PDF resumes)")
        return "\n".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages)
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            sys.exit("pip install python-docx  (needed to read Word resumes)")
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext in (".txt", ".md", ".markdown", ".text"):
        raw = path.read_bytes()
        try:
            return raw.decode("utf-8-sig")
        except UnicodeDecodeError:  # saved by Windows Notepad / Word as "ANSI": dashes and bullets are cp1252
            return raw.decode("cp1252", errors="replace")
    sys.exit(f"Unsupported resume type {ext}: use PDF, DOCX, TXT or MD.")


def guess_contact(text):
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone = re.search(r"(?:\+1[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}", text)
    linkedin = re.search(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w-]+/?", text, re.I)
    first_line = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")
    name = first_line if len(first_line.split()) <= 5 and not re.search(r"\d|@", first_line) else ""
    return {"name": name, "email": email.group(0) if email else "",
            "phone": phone.group(0) if phone else "", "linkedin": linkedin.group(0) if linkedin else ""}


def guess_years(text):
    """Years of experience from date ranges like '2011 - 2016' / 'Jan 2019 – Present'."""
    this_year = dt.date.today().year
    spans = []
    month = r"(?:(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s*|\d{1,2}/)?"
    for m in re.finditer(rf"((?:19|20)\d{{2}})\s*(?:-|–|—|to)\s*{month}((?:19|20)\d{{2}}|present|current|now)", text, re.I):
        start = int(m.group(1))
        end = this_year if not m.group(2)[0].isdigit() else int(m.group(2))
        if 1970 <= start <= end <= this_year:
            spans.append((start, end))
    if not spans:
        return None
    # merge overlaps so concurrent roles aren't double counted
    spans.sort()
    total, (cs, ce) = 0, spans[0]
    for s, e in spans[1:]:
        if s <= ce:
            ce = max(ce, e)
        else:
            total += ce - cs
            cs, ce = s, e
    total += ce - cs
    return total or None


def draft_candidate(text, resume_path):
    contact = guess_contact(text)
    return {
        "name": contact.pop("name"),
        "contact": contact,
        "base_resume": str(resume_path),
        "years_experience": guess_years(text),
        "headline": "",
        "target_titles": [],
        "past_titles": [],
        "target_levels": [],
        "skills": {"core": [], "other": []},
        "industries": [],
        "summary": "",
        "facts_files": [],
    }


# ----------------------------- candidate.yaml -------------------------------

class Candidate:
    def __init__(self, data: dict):
        self.raw = data
        self.name = data.get("name", "")
        self.years = data.get("years_experience")
        self.target_titles = [str(t) for t in data.get("target_titles", []) or []]
        self.past_titles = [str(t) for t in data.get("past_titles", []) or []]
        lv = [str(x).lower() for x in data.get("target_levels", []) or []]
        self.target_levels = [LEVEL_INDEX[x] for x in lv if x in LEVEL_INDEX]
        sk = data.get("skills") or {}
        if isinstance(sk, list):
            sk = {"core": sk}
        self.skills = {}
        for s in sk.get("other", []) or []:
            self.skills[str(s).lower()] = 1.0
        for s in sk.get("core", []) or []:
            self.skills[str(s).lower()] = 2.0
        self.industries = [str(x).lower() for x in data.get("industries", []) or []]

    @property
    def usable(self):
        return bool(self.skills or self.target_titles or self.past_titles)


def load_candidate(path: Path):
    path = Path(path)
    if not path.exists():
        return None
    try:
        with open(path, encoding="utf-8") as fh:
            data = yaml.safe_load(fh) or {}
    except yaml.YAMLError as e:
        sys.exit(f"{path.name} isn't valid YAML: {e}")
    from .validate import validate_candidate
    errors, _ = validate_candidate(data)
    if errors:
        sys.exit(f"{path.name} has problems:\n" + "\n".join(f"  - {e}" for e in errors))
    return Candidate(data)


def llm_fill(candidate: dict, text: str, model: str) -> dict:
    """Ask Claude to fill titles / levels / skills from the resume text. Never invents contact info."""
    import json

    import anthropic

    prompt = (
        "Here is a resume. Return ONLY JSON with these keys, using only what the resume supports:\n"
        '{"headline": str, "past_titles": [str], "target_titles": [str, 5-10 realistic next-role search titles],\n'
        f' "target_levels": [subset of {LEVELS}], "skills": {{"core": [10-25 short skill phrases a job posting would use],'
        ' "other": [str]}, "industries": [str], "summary": "3 plain sentences"}\n\nRESUME:\n' + text[:15000])
    msg = anthropic.Anthropic().messages.create(model=model, max_tokens=2000,
                                                messages=[{"role": "user", "content": prompt}])
    data = json.loads(re.sub(r"```json|```", "", msg.content[0].text).strip())
    for k in ("headline", "past_titles", "target_titles", "target_levels", "skills", "industries", "summary"):
        if data.get(k):
            candidate[k] = data[k]
    return candidate
